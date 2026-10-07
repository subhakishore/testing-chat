

import re
HW_AUDIT_LOGIC = [
  {
    "param": "COMPUTE ALL COMMANDS",
    "logic": ""
  },
  {
    "param": "All commands",
    "logic": "Old Logic"
  },
  {
    "param": "Uptime",
    "logic": "Not Required"
  },
  {
    "param": "Last Reboot count",
    "logic": "Not Required"
  },
  {
    "param": "System Online since Last Reboot",
    "logic": "Not Required"
  },
  {
    "param": "Last Reboot List Dates",
    "logic": "Not Required"
  },
  {
    "param": "Product Number",
    "logic": "If found NOK or No Data  , need to raise"
  },
  {
    "param": "Serial Number",
    "logic": "If found NOK or No Data  , need to raise"
  },
  {
    "param": "BMC IP Address",
    "logic": "If found NOK or No Data  , need to raise"
  },
  {
    "param": "BMC MAC Address",
    "logic": "If found NOK or No Data  , need to raise"
  },
  {
    "param": "Critical Sensor Data Record",
    "logic": "If found NOK , need to raise"
  },
  {
    "param": "Non-Responsive Sensor Data Record",
    "logic": "If found NOK , need to raise"
  },
  {
    "param": "CMOS Sensor Status from BMC SEL Logs",
    "logic": "Not Required"
  },
  {
    "param": "CMOS Current Voltage",
    "logic": "Not Required"
  },
  {
    "param": "CMOS Last 30 Days Asserted and Deasserted events count from BMC SEL Logs",
    "logic": "Not Required"
  },
  {
    "param": "Disk error from BMC SEL Logs",
    "logic": "If found NOK , need to raise"
  },
  {
    "param": "Uncorrectable ECC error from BMC SEL Logs",
    "logic": "If found NOK , need to raise"
  },
  {
    "param": "Uncorrectable ECC Dates from BMC SEL Logs",
    "logic": "Not Required"
  },
  {
    "param": "Correctable ECC error from BMC SEL Logs",
    "logic": "Not Required"
  },
  {
    "param": "Correctable ECC logging limit reached from BMC SEL Logs",
    "logic": "If found NOK , need to raise"
  },
  {
    "param": "Memory Device Disabled error from BMC SEL Logs",
    "logic": "Not Required"
  },
  {
    "param": "CPU/Processor error from BMC SEL Logs",
    "logic": "If found NOK , need to raise"
  },
  {
    "param": "CPU/Processor error Dates from BMC SEL Logs",
    "logic": "If found NOK , need to raise"
  },
  {
    "param": "Number of DIMM Modules",
    "logic": "If found NOK and value is less than 12 , need to raise"
  },
  {
    "param": "Individual DIMM Size",
    "logic": "If found NOK or No data and value is less than 16 GB , need to raise"
  },
  {
    "param": "DIMM Total Size",
    "logic": "If found NOK and value is less than 187 , need to raise"
  },
  {
    "param": "Each DIMM Serial,Manufacture,Speed & Slot",
    "logic": "If found NOK or No Data , need to raise"
  },
  {
    "param": "Kernel Component Memory CE Error count from /var/log/messages",
    "logic": "Not Required"
  },
  {
    "param": "Kernel Component Memory CE Error count from dmesg -T",
    "logic": "Not Required"
  },
  {
    "param": "Kernel Component Memory CE Error count from /sys/devices/system/edac",
    "logic": "Not Required"
  },
  {
    "param": "Faulty DIMM Slot Location",
    "logic": "If found NOK , need to raise"
  },
  {
    "param": "BMC BIOS Version",
    "logic": "If found No Data  , need to raise"
  },
  {
    "param": "BMC Firmware Version",
    "logic": "If found No Data  , need to raise"
  },
  {
    "param": "CPU Model",
    "logic": "Not Required"
  },
  {
    "param": "CPU Frequency Scaling",
    "logic": "if value is less than 64 for compute servers and less than 48 for storagebm, need to raise"
  },
  {
    "param": "Nic Interfaces list",
    "logic": "Minimun count  should be 4 , If  less than 4 , need to raise"
  },
  {
    "param": "Nic Interfaces link status",
    "logic": "If found NOK , need to raise"
  },
  {
    "param": "Nic Interfaces Firmware version",
    "logic": "If found No Data  , need to raise"
  },
  {
    "param": "Nic Interfaces Driver version",
    "logic": "If found No Data  , need to raise"
  },
  {
    "param": "Compute/Master Hard Disk Labels",
    "logic": "If found NOK , need to raise"
  },
  {
    "param": "Compute/Master Hard Disk Size",
    "logic": "If size is less than 960 GB ,sda and sdb are in different size ,or No Data found  , need to raise"
  },
  {
    "param": "Compute/Master Hard Disk Health",
    "logic": "Value should be PASSED or OK, If not found , need to raise"
  },
  {
    "param": "Compute/Master Hard Disk Errors / Total Uncorrected Errors",
    "logic": "Value should be No Errors Logged or OK, If not found , need to raise"
  },
  {
    "param": "Compute/Master Hard Disk Serial Number",
    "logic": "If No Data found  , need to raise"
  },
  {
    "param": "Compute/Master Hard Disk Model Number",
    "logic": "If No Data found  , need to raise"
  },
  {
    "param": "Compute/Master Hard Disk Firmware Version",
    "logic": "Not Required"
  },
  {
    "param": "Compute/Master Hard Disk Media_Wearout_Indicator",
    "logic": "If  Wear_Leveling_Count VALUE is less or equal to 5 , WORST is less or equal to 5 , need to raise"
  },
  {
    "param": "Compute/Master Hard Disk Erase_Fail_Count",
    "logic": "Not Required"
  },
  {
    "param": "Redfish Version Status",
    "logic": "Not Required"
  },
  {
    "param": "Redfish SessionTimeout",
    "logic": "if value is less 900, need to raise"
  },
  {
    "param": "Pwr and Perf Profile (OEMPAPP000)",
    "logic": "Not Required"
  },
  {
    "param": "Boot mode select (FBO001)",
    "logic": "Not Required"
  },
  {
    "param": "Configure SATA as (PCHS002)",
    "logic": "Not Required"
  },
  {
    "param": "SR-IOV Support (PCIS007)",
    "logic": "Not Required"
  },
  {
    "param": "Ipv4 PXE Support (NWSK001)",
    "logic": "Not Required"
  },
  {
    "param": "UEFI - Boot Option #1 (FBO201)",
    "logic": "Not Required"
  },
  {
    "param": "UEFI - Boot Option #2 (FBO202)",
    "logic": "Not Required"
  },
  {
    "param": "UEFI - Boot Option #3 (FBO203)",
    "logic": "Not Required"
  },
  {
    "param": "UEFI - Boot Option #4 (FBO204)",
    "logic": "Not Required"
  },
  {
    "param": "UEFI - Boot Option #5 (FBO205)",
    "logic": "Not Required"
  },
  {
    "param": "UEFI - Boot Option #6 (FBO206)",
    "logic": "Not Required"
  },
  {
    "param": "UEFI - Boot Option #7 (FBO207)",
    "logic": "Not Required"
  },
  {
    "param": "STORAGE ALL COMMANDS",
    "logic": ""
  },
  {
    "param": "All commands",
    "logic": "Logic"
  },
  {
    "param": "Storage NVME Hard Disk Labels",
    "logic": "If found NOK or No Data, need to raise"
  },
  {
    "param": "Storage NVME Hard Disk PCI Slot",
    "logic": "If found NOK or No Data, need to raise"
  },
  {
    "param": "Storage NVME Hard Disk Size",
    "logic": "If found NOK or No Data, need to raise"
  },
  {
    "param": "Storage NVME Hard Disk Health",
    "logic": "If found NOK or No Data, need to raise"
  },
  {
    "param": "Storage NVME Hard Disk Errors / Total Uncorrected Errors",
    "logic": "If found NOK or No Data, need to raise"
  },
  {
    "param": "Storage NVME Hard Disk Serial Number",
    "logic": "If found NOK or No Data, need to raise"
  },
  {
    "param": "Storage NVME Hard Disk Model Number",
    "logic": "If found NOK or No Data, need to raise"
  },
  {
    "param": "Storage NVME Hard Disk Firmware Version",
    "logic": "If found NOK or No Data, need to raise"
  },
  {
    "param": "Bad Block Count",
    "logic": "If found NOK or No Data, need to raise"
  },
  {
    "param": "CEPH Health Detail",
    "logic": "If found NOK , need to raise"
  },
  {
    "param": "CEPH list no of OSD's",
    "logic": "If found NOK and osd count should be 45 and all osd should be up,if not found  need to raise"
  },
  {
    "param": "Storage RAID Controller Model Number",
    "logic": "If found NOK , need to raise"
  },
  {
    "param": "JBOD Model Number",
    "logic": "If found NOK , need to raise"
  },
  {
    "param": "JBOD Firmware Build Package",
    "logic": "If found NOK , need to raise"
  },
  {
    "param": "JBOD Enclosure Disks Information for Storage",
    "logic": ""
  },
  {
    "param": "JBOD DISK INFORMATION",
    "logic": ""
  },
  {
    "param": "EID:SLT",
    "logic": "if Found NOK and total disk count should be  15 ,if not found , need to raise"
  },
  {
    "param": "DID",
    "logic": "Not Required"
  },
  {
    "param": "DISK STATE",
    "logic": "value should be  Onln-OK ,if not found , need to raise"
  },
  {
    "param": "DISK GROUP (DG)",
    "logic": "Not Required"
  },
  {
    "param": "DG/VD",
    "logic": "Not Required"
  },
  {
    "param": "VIRTUAL DISK (VD)",
    "logic": "Not Required"
  },
  {
    "param": "DISK LABEL",
    "logic": "Not Required"
  },
  {
    "param": "RAID TYPE",
    "logic": "Not Required"
  },
  {
    "param": "RAID STATE",
    "logic": "Not Required"
  },
  {
    "param": "RAID DISK ACCESS MODE",
    "logic": "Not Required"
  },
  {
    "param": "RAID DISK CACHE MODE",
    "logic": "Not Required"
  },
  {
    "param": "RAID DISK STRIP SIZE",
    "logic": "Not Required"
  },
  {
    "param": "DISK CAPACITY",
    "logic": "If found NOK and Disk size less than  480 GB , need to raise"
  },
  {
    "param": "DISK SMART HEALTH STATUS",
    "logic": "If found NOK , need to raise"
  },
  {
    "param": "ERROR STATUS/TOTAL UNCORRECTED ERRORS",
    "logic": "If found NOK , need to raise"
  },
  {
    "param": "DISK SERIAL NUMBER",
    "logic": "If found NOK or No Data  , need to raise"
  },
  {
    "param": "DISK MODEL NUMBER",
    "logic": "If found NOK or No Data  , need to raise"
  },
  {
    "param": "DISK FIRMWARE VERSION",
    "logic": "If found NOK or No Data  , need to raise"
  },
  {
    "param": "MEDIA WEAROUT INDICATOR",
    "logic": "If found NOK and  VALUE is less or equal to 5 , WORST is less or equal to 5 , need to raise"
  },
  {
    "param": "ERASE FAIL COUNT",
    "logic": "If found NOK and  VALUE is less or equal to 5 , WORST is less or equal to 5 , need to raise"
  },
  {
    "param": "OSD DETAILS",
    "logic": ""
  },
  {
    "param": "OSD ID",
    "logic": "OSD Count should be 13 , if found less than 13 , need to raise"
  },
  {
    "param": "OSD STATUS",
    "logic": "OSD Count should be 13 and all should be up-OK,if found less than 13 or Down-NOK, need to raise"
  },
  {
    "param": "OSD DEVICE LABEL",
    "logic": "Not Required"
  },
  {
    "param": "RAID TYPE",
    "logic": "Not Required"
  },
  {
    "param": "RAID STATE",
    "logic": "Not Required"
  },
  {
    "param": "EID:SLT",
    "logic": "Not Required"
  },
  {
    "param": "DID",
    "logic": "Not Required"
  },
  {
    "param": "DISK STATE",
    "logic": "OSD Count should be 13 and all should be Onln-OK, if not found  , need to raise"
  },
  {
    "param": "DG/VD",
    "logic": "Not Required"
  },
  {
    "param": "DISK GROUP (DG)",
    "logic": "Not Required"
  },
  {
    "param": "VIRTUAL DISK (VD)",
    "logic": "Not Required"
  },
  {
    "param": "DISK CAPACITY",
    "logic": "Not Required"
  },
  {
    "param": "DISK SMART HEALTH STATUS",
    "logic": "If found NOK , need to raise"
  },
  {
    "param": "FAST OSD DETAILS",
    "logic": ""
  },
  {
    "param": "FAST OSD ID",
    "logic": "Fast OSD Count should be 2 , if found less than 2 and Found NOK , need to raise"
  },
  {
    "param": "FAST OSD STATUS",
    "logic": "Fast OSD Count should be 2 , if found less than 2 and Found NOK , need to raise"
  },
  {
    "param": "FAST DEVICE LABEL",
    "logic": "Not required"
  },
  {
    "param": "NVME DISK PCI SLOT",
    "logic": "Not required"
  },
  {
    "param": "NVME DISK SIZE",
    "logic": "If Found No Data , need to raise"
  },
  {
    "param": "NVME DISK HEALTH",
    "logic": "Value should be PASSED, If not found , need to raise"
  },
  {
    "param": "NVME DISK ERRORS",
    "logic": "Value should be No Errors Logged, If not found , need to raise"
  },
  {
    "param": "NVME DISK SERIAL NUMBER",
    "logic": "If found NOK or No Data , need to raise"
  },
  {
    "param": "NVME DISK MODEL NUMBER",
    "logic": "If found NOK or No Data , need to raise"
  },
  {
    "param": "NVME DISK FIRMWARE VERSION",
    "logic": "If found NOK or No Data , need to raise"
  }
]
CMOS_THRESHOLD = 3.50

def evaluate_hw_param(param, value_str, host_name):
    val = str(value_str or "").strip()
    lower = val.lower()
    logic_entry = next((r for r in HW_AUDIT_LOGIC if r["param"].lower() == param.lower()), None)
    logic = logic_entry["logic"] if logic_entry else ""
    logic_lower = logic.lower()
    
    if not logic or "not required" in logic_lower or logic.strip() == "":
        # Keep CMOS check even if logic says Not Required per user request
        if param != "CMOS Current Voltage":
            return {"ticket": False, "reason": "Not Required"}
    
    has_nok = "nok" in lower
    has_no_data = "no data" in lower
    
    def extract_number(s):
        m = re.search(r"([0-9]+\.?[0-9]*)", s)
        return float(m.group(1)) if m else None
    def extract_gb(s):
        m = re.search(r"([0-9]+)\s*GB", s, re.I)
        return float(m.group(1)) if m else None
    
    if "if found nok or no data" in logic_lower:
        if has_nok or has_no_data:
            return {"ticket": True, "reason": f"{'NOK' if has_nok else ''} {'No Data' if has_no_data else ''}".strip()}
    elif "if found nok" in logic_lower:
        if has_nok:
            return {"ticket": True, "reason": "Found NOK"}
    elif "no data" in logic_lower and "need to raise" in logic_lower:
        if has_no_data or ("nok" in logic_lower and has_nok):
            if has_no_data:
                return {"ticket": True, "reason": "No Data"}
            if has_nok:
                return {"ticket": True, "reason": "NOK"}
    
    if param == "Number of DIMM Modules":
        num = extract_number(val)
        if has_nok or (num is not None and num < 12):
            return {"ticket": True, "reason": f"DIMM count {num} <12 or NOK"}
    if param == "Individual DIMM Size":
        gb = extract_gb(val)
        if has_nok or has_no_data or (gb is not None and gb < 16):
            return {"ticket": True, "reason": f"DIMM {gb}GB <16"}
    if param == "DIMM Total Size":
        num = extract_number(val)
        if has_nok or (num is not None and num < 187):
            return {"ticket": True, "reason": f"Total {num}<187"}
    if param == "CPU Frequency Scaling":
        num = extract_number(val)
        is_storage = "storagebm" in host_name.lower()
        threshold = 48 if is_storage else 64
        if num is not None and num < threshold:
            return {"ticket": True, "reason": f"CPU scaling {num} <{threshold}"}
    if param == "Nic Interfaces list":
        count = len([s for s in val.split(",") if s.strip()])
        if count>0 and count < 4:
            return {"ticket": True, "reason": f"NIC count {count}<4"}
    if param == "Compute/Master Hard Disk Size":
        sizes = [float(m.group(1)) for m in re.finditer(r"=>\s*([0-9]+)\s*GB", val, re.I)]
        if has_no_data:
            return {"ticket": True, "reason": "No Data"}
        if sizes:
            if any(s < 960 for s in sizes):
                return {"ticket": True, "reason": f"Disk <960GB {sizes}"}
            if len(sizes)>=2 and sizes[0] != sizes[1]:
                return {"ticket": True, "reason": f"sda/sdb mismatch {sizes[0]} vs {sizes[1]}"}
    if param == "Compute/Master Hard Disk Health":
        if "passed" not in lower and "ok" not in lower:
            return {"ticket": True, "reason": "Health not PASSED/OK"}
    if param == "Compute/Master Hard Disk Errors / Total Uncorrected Errors":
        if "no errors" not in lower and "ok" not in lower:
            return {"ticket": True, "reason": "Errors found"}
    if "media_wearout_indicator" in param.lower() or param == "MEDIA WEAROUT INDICATOR":
        vm = re.search(r"VALUE\s*([0-9]+)", val, re.I)
        wm = re.search(r"WORST\s*([0-9]+)", val, re.I)
        v = int(vm.group(1)) if vm else None
        w = int(wm.group(1)) if wm else None
        if (v is not None and v <=5) or (w is not None and w <=5):
            return {"ticket": True, "reason": f"Wearout VALUE {v} WORST {w} <=5"}
    if param == "Redfish SessionTimeout":
        num = extract_number(val)
        if num is not None and num < 900:
            return {"ticket": True, "reason": f"Timeout {num}<900"}
    if param == "DISK CAPACITY":
        gb = extract_gb(val) or extract_number(val)
        if has_nok or (gb is not None and gb < 480):
            return {"ticket": True, "reason": f"Capacity {gb}<480 or NOK"}
    if param == "CMOS Current Voltage":
        m = re.search(r"([0-9]+\.?[0-9]*)\s*volts", val, re.I)
        if m:
            volt = float(m.group(1))
            if volt > CMOS_THRESHOLD:
                return {"ticket": True, "reason": f"CMOS {volt}V > {CMOS_THRESHOLD}V"}
    # Generic less than
    lt_match = re.search(r"less than\s*([0-9]+)", logic_lower)
    if lt_match:
        threshold = float(lt_match.group(1))
        gb = extract_gb(val)
        num = extract_number(val)
        if "gb" in logic_lower and gb is not None and gb < threshold:
            return {"ticket": True, "reason": f"{gb}GB <{threshold}GB"}
        if num is not None and num < threshold:
            return {"ticket": True, "reason": f"{num}<{threshold}"}
    
    if has_nok and "nok" in logic_lower:
        return {"ticket": True, "reason": "NOK per logic"}
    if has_no_data and "no data" in logic_lower:
        return {"ticket": True, "reason": "No Data per logic"}
    
    return {"ticket": False, "reason": "OK"}

# app.py - CHATBOT LIKE ME - FINAL GENERIC - NO HARDCODING
from fastapi import FastAPI, UploadFile, File
from fastapi.middleware.cors import CORSMiddleware
from fastapi.staticfiles import StaticFiles
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
from langchain_ollama import ChatOllama
from langchain_huggingface import HuggingFaceEmbeddings
from langchain_chroma import Chroma
from langchain_core.prompts import ChatPromptTemplate
import os, traceback
from pathlib import Path
from collections import Counter
import difflib, io
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Border, Side
try:
    from rapidfuzz import fuzz
    HAS_RAPIDFUZZ = True
except:
    HAS_RAPIDFUZZ = False
    class DummyFuzz:
        def partial_ratio(self, a, b): return 50
    fuzz = DummyFuzz()

BASE_DIR = Path(__file__).parent
DATA_ROOT = BASE_DIR / "data"
DB_PATH = str(BASE_DIR.parent / "vector_db")
IMAGE_ROOT = BASE_DIR / "extracted_images"
LLM_MODEL = "llama3.2"

app = FastAPI(title="TASS Chatbot Backend")
app.add_middleware(CORSMiddleware, allow_origins=["*"], allow_credentials=False, allow_methods=["*"], allow_headers=["*"])
os.makedirs(IMAGE_ROOT, exist_ok=True)
for t in ["hotline","ran","tc30","troubleshoot","comparison","audit"]:
    (IMAGE_ROOT / t).mkdir(parents=True, exist_ok=True)
app.mount("/images", StaticFiles(directory=str(IMAGE_ROOT)), name="images")

embeddings = HuggingFaceEmbeddings(model_name="sentence-transformers/all-MiniLM-L6-v2")
try:
    vector_db = Chroma(persist_directory=DB_PATH, embedding_function=embeddings)
    print(f"DB loaded: {vector_db._collection.count()} docs")
except Exception as e:
    vector_db = Chroma(persist_directory=DB_PATH, embedding_function=embeddings)

llm = ChatOllama(model=LLM_MODEL, temperature=0.2)

class QueryRequest(BaseModel):
    question: str
    topic: str = "hotline"

def normalize_topic(t: str) -> str:
    t = t.lower().strip()
    mapping = {"helpline":"hotline","tsh":"troubleshoot","ran_support":"ran","tc30_support":"tc30"}
    return mapping.get(t,t)

# CHATBOT LIKE ME - ONE GENERIC PROMPT - NO TSH NAMES
prompt_template = ChatPromptTemplate.from_template("""
You are TASS, a friendly and helpful NOKIA support chatbot like Meta AI.

Your job:
- Answer ONLY from the Context below.
- Be conversational, friendly, like a real chatbot.
- If question is a TSH issue, give clear steps with headings using <b>, <ul><li>.
- If question asks for definition/abbreviation/full form, give definition only.
- If answer is NOT in Context, say politely: "I couldn't find this in the knowledge base for {topic}. It seems this issue is not documented in the provided files."
- Do NOT invent steps, do NOT reuse other TSH, do NOT hallucinate.
- Keep alarm codes, commands, versions exactly as in Context.

Topic: {topic}
Context:
{context}

Question: {question}

Answer in friendly HTML format allowed: <b>, <br>, <ul><li>, <table>, <pre>. Do NOT wrap in ```html.
If context mentions image/diagram, mention it.

Answer:
""")

@app.post("/ask")
async def ask_question(req: QueryRequest):
    try:
        topic = normalize_topic(req.topic)
        q_lower = req.question.lower()
        q_keywords = [w for w in q_lower.split() if len(w) > 2]
        print(f"\nQuery: {req.question} | Topic: {topic}")

        all_vector_results = vector_db.similarity_search_with_score(req.question, k=10)
        strict_filtered = [(d, s) for d, s in all_vector_results if d.metadata.get("topic","").lower() == topic]

        # Generic fallback: if nothing in this topic, search all topics
        if not strict_filtered:
            strict_filtered = all_vector_results[:10]

        if not strict_filtered:
            return {"answer": f"I couldn't find anything for '{req.question}' in topic <b>{topic}</b>.", "sources": [], "images": []}

        reranked=[]
        for doc,score in strict_filtered:
            try: s=float(score)
            except: s=0.5
            keyword_score = sum(1 for kw in q_keywords if kw in doc.page_content.lower())
            fuzzy_score = fuzz.partial_ratio(q_lower, doc.page_content.lower()) if HAS_RAPIDFUZZ else (50 if q_lower in doc.page_content.lower() else 0)
            sim = 1.0/(1.0+s)  # Fix: proper similarity, not 1-s
            combined = sim*0.5 + (keyword_score/max(len(q_keywords),1))*0.25 + (fuzzy_score/100)*0.25
            reranked.append((doc,score,combined,keyword_score,fuzzy_score))

        reranked.sort(key=lambda x: x[2], reverse=True)
        top_docs_with_scores = reranked[:2]
        top_docs = [doc for doc,_,_,_,_ in top_docs_with_scores]

        # GENERIC CHECK - NO HARDCODING ANY TSH
        if top_docs_with_scores:
            best_kw = top_docs_with_scores[0][3]
            best_fuzzy = top_docs_with_scores[0][4]
            kw_cov = best_kw / max(len(q_keywords),1)
            print(f"  BEST kw={best_kw}/{len(q_keywords)}={kw_cov:.2f} fuzzy={best_fuzzy} combined={top_docs_with_scores[0][2]:.3f}")
            # If less than 35% keywords match AND fuzzy low => not in docs
            if kw_cov < 0.35 and best_fuzzy < 65:
                return {
                    "answer": f"I couldn't find this in the knowledge base for <b>{topic}</b>. It seems this issue <b>'{req.question}'</b> is not documented in the provided files.",
                    "sources": [],
                    "images": []
                }

        context_text = "\n\n---\n\n".join([d.page_content for d in top_docs])
        chain = prompt_template | llm
        response = chain.invoke({"context": context_text, "question": req.question, "topic": topic})
        final_answer = response.content

        sources, seen = [], set()
        for doc,_,_,_,_ in top_docs_with_scores[:2]:
            src = doc.metadata.get("source","unknown")
            if src not in seen:
                sources.append(src)
                seen.add(src)
        return {"answer": final_answer, "sources": sources, "images": []}

    except Exception as e:
        traceback.print_exc()
        return {"answer": f"Error: {str(e)}", "sources": [], "images": []}

@app.post("/query")
async def query_alias(req: QueryRequest):
    return await ask_question(req)

# Keep compare endpoints
GREEN_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
RED_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
YELLOW_FILL = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
GREEN_FONT = Font(color="006100", bold=True)
RED_FONT = Font(color="9C0006", bold=True)
YELLOW_FONT = Font(color="9C6500", bold=True)
BOLD_FONT = Font(bold=True)
THIN_BORDER = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))

@app.post("/compare")
async def compare_and_download_excel(pre_file: UploadFile = File(...), post_file: UploadFile = File(...)):
    pre_content = (await pre_file.read()).decode('utf-8', errors='ignore')
    post_content = (await post_file.read()).decode('utf-8', errors='ignore')
    pre_lines = pre_content.splitlines(); post_lines = post_content.splitlines()
    matcher = difflib.SequenceMatcher(None, pre_lines, post_lines)
    wb = Workbook(); ws = wb.active; ws.title = "Comparison"
    ws.append(["Line No", "Status", "Pre-Check Content", "Post-Check Content"])
    er=2
    for tag,i1,i2,j1,j2 in matcher.get_opcodes():
        if tag=='equal':
            for k in range(i1,i2): ws.append([er-1, "UNCHANGED", pre_lines[k], post_lines[k]]); er+=1
        elif tag=='delete':
            for k in range(i1,i2): ws.append([er-1, "REMOVED", pre_lines[k], ""]); er+=1
        elif tag=='insert':
            for k in range(j1,j2): ws.append([er-1, "ADDED", "", post_lines[k]]); er+=1
        else:
            ml=max(i2-i1,j2-j1)
            for k in range(ml):
                pl=pre_lines[i1+k] if i1+k<i2 else ""; pol=post_lines[j1+k] if j1+k<j2 else ""
                ws.append([er-1, "CHANGED", pl, pol]); er+=1
    buf=io.BytesIO(); wb.save(buf); buf.seek(0)
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": "attachment; filename=Pre_Post_Comparison.xlsx"})

@app.post("/compare-json")
async def compare_json_preview(pre_file: UploadFile = File(...), post_file: UploadFile = File(...)):
    pre = (await pre_file.read()).decode('utf-8', errors='ignore').splitlines()
    post = (await post_file.read()).decode('utf-8', errors='ignore').splitlines()
    m = difflib.SequenceMatcher(None, pre, post)
    res=[]
    for tag,i1,i2,j1,j2 in m.get_opcodes():
        if tag=='equal':
            for k in range(i1,i2): res.append({"status":"UNCHANGED","pre":pre[k],"post":post[k]})
        elif tag=='delete':
            for k in range(i1,i2): res.append({"status":"REMOVED","pre":pre[k],"post":""})
        elif tag=='insert':
            for k in range(j1,j2): res.append({"status":"ADDED","pre":"","post":post[k]})
        else:
            ml=max(i2-i1,j2-j1)
            for k in range(ml): res.append({"status":"CHANGED","pre":pre[i1+k] if i1+k<i2 else "","post":post[j1+k] if j1+k<j2 else ""})
    return {"summary":f"{len(res)} differences","data":res[:500]}


@app.post("/compare-inline")
async def compare_inline(pre_file: UploadFile = File(...), post_file: UploadFile = File(...)):
    pre_content = (await pre_file.read()).decode('utf-8', errors='ignore')
    post_content = (await post_file.read()).decode('utf-8', errors='ignore')
    pre_lines = pre_content.splitlines()
    post_lines = post_content.splitlines()
    matcher = difflib.SequenceMatcher(None, pre_lines, post_lines)
    html = ['<div style="font-family:monospace; font-size:12px;">']
    html.append('<div style="display:flex; background:#0f172a; color:white; padding:8px; font-weight:700;"><div style="flex:1;">Pre</div><div style="flex:1;">Post</div></div>')
    for tag,i1,i2,j1,j2 in matcher.get_opcodes():
        if tag=='equal':
            for k in range(i1,i2):
                html.append(f'<div style="display:flex; border-bottom:1px solid #f1f5f9;"><div style="flex:1; padding:4px 8px; background:#f8fafc;">{pre_lines[k][:500]}</div><div style="flex:1; padding:4px 8px; background:#f8fafc;">{post_lines[k][:500]}</div></div>')
        elif tag=='delete':
            for k in range(i1,i2):
                html.append(f'<div style="display:flex; border-bottom:1px solid #f1f5f9;"><div style="flex:1; padding:4px 8px; background:#fee2e2;">- {pre_lines[k][:500]}</div><div style="flex:1; padding:4px 8px; background:#f1f5f9;"></div></div>')
        elif tag=='insert':
            for k in range(j1,j2):
                html.append(f'<div style="display:flex; border-bottom:1px solid #f1f5f9;"><div style="flex:1; padding:4px 8px; background:#f1f5f9;"></div><div style="flex:1; padding:4px 8px; background:#dcfce7;">+ {post_lines[k][:500]}</div></div>')
        else:
            ml=max(i2-i1,j2-j1)
            for k in range(ml):
                pl=pre_lines[i1+k] if i1+k<i2 else ""
                pol=post_lines[j1+k] if j1+k<j2 else ""
                html.append(f'<div style="display:flex; border-bottom:1px solid #f1f5f9;"><div style="flex:1; padding:4px 8px; background:#fef3c7;">~ {pl[:500]}</div><div style="flex:1; padding:4px 8px; background:#fef3c7;">~ {pol[:500]}</div></div>')
    html.append('</div>')
    from fastapi.responses import HTMLResponse
    return HTMLResponse(content="".join(html))


@app.post("/audit/analyze")
async def audit_analyze(file: UploadFile = File(...)):
    import re, csv, io
    content = await file.read()
    filename = file.filename
    text = content.decode('utf-8', errors='ignore')
    lines = text.splitlines()
    reader = csv.reader(lines)
    rows = list(reader)
    if len(rows) < 3:
        try:
            import pandas as pd
            df = pd.read_excel(io.BytesIO(content))
            rows = [df.columns.tolist()] + df.values.tolist()
        except Exception as e:
            return {"error": f"Could not parse file: {e}"}
    
    # Find hosts row
    hosts_row = None
    for r in rows[:5]:
        if any('mnnaf' in str(x) or 'All commands' in str(x) for x in r):
            if 'All commands' in str(r[0]):
                hosts_row = r
                break
    if not hosts_row and len(rows)>=3:
        hosts_row = rows[2]
    
    hosts = []
    if hosts_row:
        for idx, h in enumerate(hosts_row):
            hn = str(h).strip()
            if hn and hn != "All commands" and "COMPUTE" not in hn and len(hn)>2:
                hosts.append({"col": idx, "name": hn})
    
    # Evaluate all params
    violations = []
    for row in rows:
        if not row or not row[0]:
            continue
        param = str(row[0]).strip()
        if not param or "COMPUTE ALL COMMANDS" in param or param == "All commands" or param == "":
            continue
        logic_entry = next((r for r in HW_AUDIT_LOGIC if r["param"].lower() == param.lower()), None)
        if not logic_entry:
            continue
        if "not required" in logic_entry["logic"].lower() and param != "CMOS Current Voltage":
            if logic_entry["logic"].strip() == "" or "not required" in logic_entry["logic"].lower():
                # Skip but keep CMOS
                if param != "CMOS Current Voltage":
                    continue
        
        for host in hosts:
            col = host["col"]
            if col >= len(row):
                continue
            value = str(row[col]).strip()
            if not value:
                continue
            result = evaluate_hw_param(param, value, host["name"])
            if result["ticket"]:
                violations.append({
                    "host": host["name"],
                    "param": param,
                    "value": value,
                    "logic": logic_entry["logic"],
                    "reason": result["reason"],
                    "ticketRequired": "YES",
                    "action": f"Raise Ticket - {param}: {result['reason']}"
                })
    
    # Create Excel with single sheet for all violations
    wb = Workbook()
    ws_main = wb.active
    ws_main.title = "All_Violations_Single_Sheet"
    ws_main.append(["Host", "Parameter", "Value", "Logic Rule", "Failure Reason", "Ticket Required", "Action", "Severity", "File", "Date"])
    for v in violations:
        severity = "Medium"
        if "critical" in v["param"].lower() or "disk error" in v["param"].lower() or "cpu" in v["param"].lower() or "dimm" in v["param"].lower():
            severity = "High"
        if "CMOS" in v["param"]:
            severity = "High"
        ws_main.append([
            v["host"], v["param"], v["value"], v["logic"], v["reason"], "YES", v["action"], severity, filename, ""
        ])
    
    # Summary
    ws_summary = wb.create_sheet("Summary")
    ws_summary.append(["HW Audit Full Logic Report"])
    ws_summary.append(["File", filename])
    ws_summary.append(["Total Hosts", len(hosts)])
    ws_summary.append(["Total Rules", len(HW_AUDIT_LOGIC)])
    ws_summary.append(["Total Violations", len(violations)])
    ws_summary.append(["Unique Hosts With Issues", len(set(v["host"] for v in violations))])
    ws_summary.append(["CMOS Threshold", f"{CMOS_THRESHOLD}V"])
    ws_summary.append([])
    ws_summary.append(["Breakdown by Parameter"])
    from collections import Counter
    cnt = Counter(v["param"] for v in violations)
    for p,c in cnt.items():
        ws_summary.append([p, c])
    
    # Full data
    ws_full = wb.create_sheet("Full_Audit_Data")
    for r in rows:
        ws_full.append(r)
    
    # Logic reference
    ws_logic = wb.create_sheet("Logic_Reference")
    ws_logic.append(["Param", "Logic Rule"])
    for r in HW_AUDIT_LOGIC:
        ws_logic.append([r["param"], r["logic"]])
    
    # Style header
    for cell in ws_main[1]:
        cell.font = BOLD_FONT
        cell.border = THIN_BORDER
    for row in ws_main.iter_rows(min_row=2):
        for cell in row:
            cell.border = THIN_BORDER
            if cell.col_idx == 6:  # Ticket Required
                if cell.value == "YES":
                    cell.fill = RED_FILL
                    cell.font = RED_FONT
    
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": f"attachment; filename=HW_Audit_Full_Report_{filename.rsplit('.',1)[0]}.xlsx"})


@app.post("/audit/cmos-excel")
async def audit_cmos_excel(file: UploadFile = File(...)):
    return await audit_analyze(file)

@app.get("/")
def health():
    try:
        count = vector_db._collection.count()
        metas = vector_db._collection.get(include=["metadatas"])["metadatas"]
        topics = Counter([m.get("topic","NO_TOPIC") for m in metas])
        return {"status":"TASS chatbot running","vector_count":count,"topics":dict(topics)}
    except Exception as e:
        return {"status":"error","error":str(e)}
